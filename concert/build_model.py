# -*- coding: utf-8 -*-
"""Бюджет концерта — один лист, строки и суммы."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

F = 'Arial'
RUB = '#,##0" ₽";[Red]-#,##0" ₽";"—"'
PCT = '0.0%'
INT = '#,##0;[Red]-#,##0;"—"'
BLUE, GREY = '0000FF', '808080'
FILL_IN = PatternFill('solid', fgColor='FFF9E6')
FILL_SEC = PatternFill('solid', fgColor='EDF2F9')
THIN = Side(style='thin', color='D0D0D0')
TOP = Side(style='thin', color='404040')
DBL = Side(style='double', color='404040')

wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Бюджет'
ws.sheet_view.showGridLines = False
for col, w in {'A': 56, 'B': 11, 'C': 15, 'D': 18, 'E': 2, 'F': 58}.items():
    ws.column_dimensions[col].width = w

r = 1
ws.cell(r, 1, 'БЮДЖЕТ КОНЦЕРТА').font = Font(name=F, size=16, bold=True)
r += 1
ws.cell(r, 1, 'Заполняйте ячейки на жёлтом фоне. Суммы и итоги считаются сами.').font = Font(
    name=F, size=9, italic=True, color=GREY)
r += 2

def head():
    global r
    for i, h in enumerate(['Статья', 'Кол-во', 'Цена / ставка', 'Сумма, ₽']):
        c = ws.cell(r, 1 + i, h)
        c.font = Font(name=F, size=9, bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='44546A')
        c.alignment = Alignment(horizontal='center' if i else 'left')
        c.border = Border(bottom=TOP)
    r += 1

def section(text):
    global r
    c = ws.cell(r, 1, text)
    c.font = Font(name=F, size=10, bold=True, color='1F3864')
    for i in range(1, 5):
        ws.cell(r, i).fill = FILL_SEC
        ws.cell(r, i).border = Border(bottom=THIN)
    r += 1

def item(label, qty, price, fmt_price=RUB, note=None, sign=-1):
    """Строка сметы: кол-во × цена = сумма."""
    global r
    ws.cell(r, 1, label).font = Font(name=F, size=10)
    ws.cell(r, 1).alignment = Alignment(indent=1)
    for col, val, fmt in ((2, qty, INT), (3, price, fmt_price)):
        c = ws.cell(r, col, val)
        c.number_format = fmt
        c.font = Font(name=F, size=10, color=BLUE)
        c.fill = FILL_IN
        c.alignment = Alignment(horizontal='center' if col == 2 else 'right')
    c = ws.cell(r, 4, f'={"-" if sign < 0 else ""}B{r}*C{r}')
    c.number_format = RUB
    c.font = Font(name=F, size=10)
    if note:
        ws.cell(r, 6, note).font = Font(name=F, size=9, italic=True, color=GREY)
    r += 1
    return r - 1

def rate(label, pct, base_ref, note=None):
    """Строка-процент от базы."""
    global r
    ws.cell(r, 1, label).font = Font(name=F, size=10)
    ws.cell(r, 1).alignment = Alignment(indent=1)
    c = ws.cell(r, 3, pct)
    c.number_format = PCT
    c.font = Font(name=F, size=10, color=BLUE)
    c.fill = FILL_IN
    c.alignment = Alignment(horizontal='right')
    d = ws.cell(r, 4, f'=-{base_ref}*C{r}')
    d.number_format = RUB
    d.font = Font(name=F, size=10)
    if note:
        ws.cell(r, 6, note).font = Font(name=F, size=9, italic=True, color=GREY)
    r += 1
    return r - 1

def total(label, formula, big=False, note=None):
    global r
    c = ws.cell(r, 1, label)
    c.font = Font(name=F, size=12 if big else 10, bold=True)
    d = ws.cell(r, 4, formula)
    d.number_format = RUB
    d.font = Font(name=F, size=12 if big else 10, bold=True)
    for i in range(1, 5):
        ws.cell(r, i).border = Border(top=DBL if big else TOP)
    if big:
        ws.row_dimensions[r].height = 22
    if note:
        ws.cell(r, 6, note).font = Font(name=F, size=9, italic=True, color=GREY)
    r += 1
    return r - 1

def blank():
    global r
    r += 1

# ---------------------------------------------------------------- билеты
head()
section('БИЛЕТЫ')
t0 = r
for lab, qty, price, note in [
    ('Ранняя пташка (early bird)',           57, 1800, 'Первые 2 недели продаж'),
    ('Стандарт (предпродажа)',               90, 2500, 'Основная квота зала'),
    ('Стандарт (на входе)',                  20, 2800, 'Цена в день концерта выше'),
    ('Фан-зона у сцены / VIP',               18, 4000, ''),
    ('Льготный (студенты)',                   7, 1500, ''),
    ('Промо, пресс, гости',                  10,    0, 'Занимают места, денег не приносят'),
]:
    item(lab, qty, price, note=note, sign=1)
t1 = r - 1
SOLD = f'SUM(B{t0}:B{t1})'
GROSS_ROW = total('1. ВАЛОВОЙ СБОР', f'=SUM(D{t0}:D{t1})')
ws.cell(GROSS_ROW, 6, f'=" Продано билетов: "&{SOLD}&" шт, из них платных "&SUM(B{t0}:B{t1-1})&" шт"')
ws.cell(GROSS_ROW, 6).font = Font(name=F, size=9, italic=True, color=GREY)
G = f'$D${GROSS_ROW}'
blank()

# ---------------------------------------------------------------- вычеты
section('ВЫЧЕТЫ ИЗ ВАЛОВОГО СБОРА')
USN  = rate('2. Налоги УСН',                  0.06, G)
COMM = rate('3. Комиссия за продажу билетов', 0.05, G, 'Считается от валового сбора')
RAO  = rate('4. Отчисления в РАО',            0.08, G)
blank()

# ---------------------------------------------------------------- техника
section('5. ТЕХНИЧЕСКИЕ РАСХОДЫ: АРЕНДА ЗВУКА И ОБЕСПЕЧЕНИЕ ПЛОЩАДКИ')
h0 = r
for lab, qty, price, note in [
    ('Аренда зала (смена)',                    1, 20000, ''),
    ('Аренда звукового комплекта PA',          1, 25000, ''),
    ('Микрофоны, радиосистемы, стойки',        1,  8000, ''),
    ('Мониторная линия / in-ear',              1, 10000, ''),
    ('Звукорежиссёр FOH',                      1, 15000, ''),
    ('Мониторный инженер',                     0, 12000, 'В клубе часто совмещается с FOH'),
    ('Световое оборудование',                  1, 15000, ''),
    ('Художник по свету / оператор',           1,  8000, ''),
    ('Бэклайн: барабаны, усилители, клавиши',  0, 20000, 'Поставьте 1, если привозной'),
    ('Экран, проекция, видеоконтент',          0, 18000, 'Поставьте 1, если есть видеоряд'),
    ('Электрика, кабели, распределение',       1,  3000, ''),
    ('Монтаж, демонтаж, погрузка',             3,  3500, ''),
    ('Охрана',                                 3,  4000, '1 человек на 75–100 гостей'),
    ('Билетный контроль и гардероб',           3,  3000, ''),
    ('Уборка до и после',                      2,  3000, ''),
    ('Дежурный медик',                         1,  5000, ''),
    ('Администратор площадки',                 1,  8000, ''),
]:
    item(lab, qty, price, note=note)
h1 = r - 1
TECH = total('Итого технических расходов', f'=SUM(D{h0}:D{h1})')
blank()

PROFIT = total('6. ПРИБЫЛЬ',
               f'=$D${GROSS_ROW}+$D${USN}+$D${COMM}+$D${RAO}+$D${TECH}',
               big=True, note='Эта сумма делится с площадкой пополам')
blank()

# ---------------------------------------------------------------- раздел
section('РАЗДЕЛ ПРИБЫЛИ')
r0 = r
ws.cell(r, 1, 'Доля площадки').font = Font(name=F, size=10)
ws.cell(r, 1).alignment = Alignment(indent=1)
c = ws.cell(r, 3, 0.50); c.number_format = PCT
c.font = Font(name=F, size=10, color=BLUE); c.fill = FILL_IN
c.alignment = Alignment(horizontal='right')
d = ws.cell(r, 4, f'=-$D${PROFIT}*C{r}'); d.number_format = RUB; d.font = Font(name=F, size=10)
VSHARE = r; r += 1
ORG = total('ПРИБЫЛЬ ОРГАНИЗАТОРА', f'=$D${PROFIT}+$D${VSHARE}', big=True,
            note='Из неё дальше вычитается всё остальное')
blank()

# ---------------------------------------------------------------- этап 2
section('РАСХОДЫ ИЗ ПРИБЫЛИ ОРГАНИЗАТОРА')
TAX2 = rate('Налоги',              0.07, f'MAX(0,$D${ORG})')
BANK = rate('Банковские комиссии', 0.01, f'MAX(0,$D${ORG})')
blank()

section('ГОНОРАРЫ АРТИСТАМ')
f0 = r
for lab, qty, price, note in [
    ('Артист / хедлайнер',       1, 25000, ''),
    ('Музыкант: гитара',         1,  8000, ''),
    ('Музыкант: бас',            1,  8000, ''),
    ('Музыкант: барабаны',       1,  8000, ''),
    ('Музыкант: клавиши',        0,  8000, ''),
    ('Разогрев / support act',   0, 10000, 'Часто выступает за билеты и промо'),
]:
    item(lab, qty, price, note=note)
f1 = r - 1
FEES = total('Итого гонорары', f'=SUM(D{f0}:D{f1})')
blank()

section('МАРКЕТИНГ И PR')
m0 = r
for lab, qty, price, note in [
    ('Таргетированная реклама VK',            1, 20000, ''),
    ('Посевы в Telegram-каналах',             1, 15000, ''),
    ('Блогеры и лидеры мнений',               1, 10000, 'Часть работает по бартеру за билеты'),
    ('Наружная реклама и расклейка афиш',     1,  5000, ''),
    ('Городские афиши: Яндекс.Афиша, KudaGo', 1,     0, 'Бесплатно, подавать за 3–4 недели'),
    ('PR: пресс-релиз, СМИ, интервью',        1,  3000, ''),
    ('Email и база подписчиков',              1,  1000, 'Самый дешёвый билет'),
    ('Дизайн, тизер, съёмка контента',        1,  6000, ''),
]:
    item(lab, qty, price, note=note)
m1 = r - 1
MKT = total('Итого маркетинг и PR', f'=SUM(D{m0}:D{m1})')
blank()

section('ПРОЧИЕ РАСХОДЫ')
o0 = r
for lab, qty, price, note in [
    ('Полиграфия: афиши, флаеры',      1, 12000, 'Печать. Размещение — в маркетинге'),
    ('Бейджи, браслеты, программки',   1,  5000, ''),
    ('Оформление сцены и зала',        0, 15000, ''),
    ('Репетиции, аренда репбазы',      2,  6000, ''),
    ('Транспорт оборудования',         2,  6000, ''),
    ('Трансфер и проживание команды',  0, 25000, 'Если команда не местная'),
    ('Кейтеринг и райдер бэкстейджа', 15,   700, 'По числу людей на площадке'),
    ('Фотограф',                       1, 12000, ''),
    ('Видеосъёмка концерта',           0, 25000, ''),
    ('Юридическое сопровождение',      1,  6000, 'Договоры с площадкой, артистами, РАО'),
]:
    item(lab, qty, price, note=note)
o1 = r - 1
RES = rate('Резерв на непредвиденное', 0.10, f'-SUM(D{o0}:D{o1})')
OTHER = total('Итого прочие расходы', f'=SUM(D{o0}:D{o1})+D{RES}')
blank()

NET = total('ЧИСТЫЙ ДОХОД ОРГАНИЗАТОРА',
            f'=$D${ORG}+$D${TAX2}+$D${BANK}+$D${FEES}+$D${MKT}+$D${OTHER}', big=True)
ws.cell(NET, 1).font = Font(name=F, size=14, bold=True)
ws.cell(NET, 4).font = Font(name=F, size=14, bold=True)
ws.row_dimensions[NET].height = 28
blank()

# ---------------------------------------------------------------- справочно
section('СПРАВОЧНО')
ws.cell(r, 1, 'Вместимость зала, чел.').font = Font(name=F, size=10)
ws.cell(r, 1).alignment = Alignment(indent=1)
c = ws.cell(r, 4, 300); c.number_format = INT
c.font = Font(name=F, size=10, color=BLUE); c.fill = FILL_IN
CAP = f'$D${r}'; r += 1

ws.cell(r, 1, 'Средняя цена платного билета').font = Font(name=F, size=10)
ws.cell(r, 1).alignment = Alignment(indent=1)
c = ws.cell(r, 4, f'=IFERROR($D${GROSS_ROW}/SUM(B{t0}:B{t1-1}),0)')
c.number_format = RUB; c.font = Font(name=F, size=10)
AVG = f'$D${r}'; r += 1

# сбор, при котором чистый доход = 0
SURV = f'(1-$C${USN}-$C${COMM}-$C${RAO})'
POST = f'(-$D${FEES}-$D${MKT}-$D${OTHER})'
ws.cell(r, 1, 'Валовой сбор для выхода в ноль').font = Font(name=F, size=10, bold=True)
ws.cell(r, 1).alignment = Alignment(indent=1)
c = ws.cell(r, 4,
    f'=IFERROR((-$D${TECH}+{POST}/((1-$C${TAX2}-$C${BANK})*(1-$C${VSHARE})))/{SURV},0)')
c.number_format = RUB; c.font = Font(name=F, size=10, bold=True)
BEG = f'$D${r}'; r += 1

ws.cell(r, 1, 'Нужно продать платных билетов').font = Font(name=F, size=10, bold=True)
ws.cell(r, 1).alignment = Alignment(indent=1)
c = ws.cell(r, 4, f'=IFERROR(ROUNDUP({BEG}/{AVG},0),0)')
c.number_format = INT; c.font = Font(name=F, size=10, bold=True)
BET = f'$D${r}'; r += 1

ws.cell(r, 1, 'Это заполняемость зала (с промо-билетами)').font = Font(name=F, size=10, bold=True)
ws.cell(r, 1).alignment = Alignment(indent=1)
c = ws.cell(r, 4, f'=IFERROR(({BET}+SUMIF(C{t0}:C{t1},0,B{t0}:B{t1}))/{CAP},0)')
c.number_format = PCT; c.font = Font(name=F, size=10, bold=True)
ws.cell(r, 6, 'Больше 100% — при этих ценах и расходах концерт не окупается').font = Font(
    name=F, size=9, italic=True, color=GREY)

ws.freeze_panes = 'A5'
ws.page_setup.orientation = 'portrait'
ws.page_setup.fitToWidth = 1
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.print_area = f'A1:D{r}'

wb.save('Бюджет_концерта.xlsx')
print(f'saved, строк: {r}, листов: {len(wb.sheetnames)}')
